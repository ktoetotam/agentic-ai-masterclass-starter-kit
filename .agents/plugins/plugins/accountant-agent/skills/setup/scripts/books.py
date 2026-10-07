#!/usr/bin/env python3
"""Accountant agent helper for the AI Realist Agentic AI Masterclass.

Python standard library, plus pypdf for PDFs and openpyxl for Excel when installed (the starter installs
both). The books are a folder:

  inbox/    a copy of every document found, named by a short id; never edited: the evidence
  text/     each document as numbered lines, so every value can be quoted with its line
  records/  one JSON file per document, written by the AI agent: the fields and the line each came from
  out/      what this script produces: register, bank lines, Stripe export, matches, review queue, workbook

Money is calculated here with Decimal, never by the AI. Nothing is sent, paid or changed anywhere: Stripe
is read with a read-only sandbox key, documents are copied, never moved. Run `python books.py --help`.
"""
from __future__ import annotations

import argparse
import base64
import csv
import datetime as dt
import email
import hashlib
import html.parser
import io
import json
import os
import re
import shutil
import ssl
import subprocess
import sys
import urllib.error
import urllib.parse
import urllib.request
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from email import policy
from pathlib import Path

VERSION = "1.0.0"
SKIP_NAMES = {"answer-key", "generator", ".git", "__pycache__", "manifest.json", "README.md", ".DS_Store"}
SECRET_NAME = re.compile(r"(^\.env|credential|secret|passw|token|private[-_]?key|\.pem$|\.key$|\.p12$)", re.I)
IMAGE = {".jpg", ".jpeg", ".png", ".heic", ".webp"}
EU = set("AT BE BG CY CZ DE DK EE ES FI FR GR HR HU IE IT LT LU LV MT NL PL PT RO SE SI SK".split())
STRIPE_API = "https://api.stripe.com/v1"
STRIPE_VERSION = "2026-09-30.endive"
QUEUE_HEADER = ["code", "item", "finding", "what_to_do", "source", "by"]
TYPES = ["sales_invoice", "cancellation", "bill", "receipt", "expense_claim", "bank_statement", "contract", "email", "note", "other"]
D = Decimal
CENT = D("0.01")


# ----------------------------------------------------------------------------------- basics

def money(v) -> Decimal | None:
    if v is None or v == "":
        return None
    try:
        s = str(v).strip().replace("EUR", "").replace("€", "").replace(" ", "")
        if re.fullmatch(r"-?[\d.]+,\d{1,2}", s):  # German 1.234,56
            s = s.replace(".", "").replace(",", ".")
        else:
            s = s.replace(",", "")
        return D(s).quantize(CENT, rounding=ROUND_HALF_UP)
    except (InvalidOperation, ValueError):
        return None


def fmt(v) -> str:
    return "" if v is None else f"{v:.2f}"


def load_books(path: str | None) -> tuple[Path, dict]:
    root = Path(path or "output/accountant")
    cfg = root / "books.json"
    if not cfg.exists():
        sys.exit(f"No books at {root}. Run: books.py init --books {root} --pack data/accountant-pack")
    return root, json.loads(cfg.read_text())


def read_csv(path: Path) -> list[dict]:
    if not path.exists():
        return []
    with path.open(newline="", encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


def write_csv(path: Path, header: list[str], rows: list[list]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(header)
        w.writerows(rows)


def tokens(text: str) -> set[str]:
    stop = {"gmbh", "ltd", "inc", "sl", "og", "ag", "the", "and", "co", "of", "demo", "online", "studio", "fictional", "&"}
    return {t for t in re.findall(r"[a-z0-9äöüß]{3,}", (text or "").lower()) if t not in stop}


def git_ignored(path: Path) -> bool | None:
    try:
        r = subprocess.run(["git", "check-ignore", "-q", str(path)], capture_output=True)
        return r.returncode == 0
    except OSError:
        return None


# ------------------------------------------------------------------------------------- init

RULES_MD = """# How these books work

Fictional workshop data unless your group decided otherwise. Edit the house rules at the end.

- `inbox/` holds a copy of every document; the originals stay where they were. Never edit `inbox/`.
- `text/<id>.txt` is each document as numbered lines. Quote values exactly as they appear there.
- `records/<id>.json` is written by the AI agent: one record per document. Every value names the line it came from.
- `out/` is written by `books.py`. Do not edit it by hand: change a record and run the command again.
- Money is calculated by `books.py` with exact decimals. The agent never adds up amounts in its head.
- The agent prepares. People approve, pay and send. Drafts are labelled DRAFT and stay in `out/`.

## Our house rules
(Copied from the company's house rules when the books were created. Add your own below.)
"""


def cmd_init(args) -> int:
    root = Path(args.books)
    pack = Path(args.pack)
    if not (pack / "drive").exists() and not (pack / "mailbox").exists():
        print(f"PROBLEM  {pack} has no drive/ or mailbox/ folder. Point --pack at data/accountant-pack or your own folder.")
        return 1
    for sub in ("inbox", "text", "records", "out"):
        (root / sub).mkdir(parents=True, exist_ok=True)
    manifest = json.loads((pack / "manifest.json").read_text()) if (pack / "manifest.json").exists() else {}
    scenario = manifest.get("scenario_date") or dt.date.today().isoformat()
    first = dt.date.fromisoformat(scenario).replace(day=1)
    month = args.month or manifest.get("month") or (first - dt.timedelta(days=1)).strftime("%Y-%m")
    cfg = root / "books.json"
    if not cfg.exists():
        company = company_from_rules(pack)
        cfg.write_text(json.dumps(dict(version=VERSION, pack=str(pack), scenario_date=scenario, month=month,
                                       company=company), indent=1) + "\n")
    rules = root / "BOOKS.md"
    if not rules.exists():
        house = next(pack.rglob("Company and house rules.md"), None)
        extra = ""
        if house:
            text = house.read_text()
            extra = text.split("## House rules", 1)[1] if "## House rules" in text else ""
        rules.write_text(RULES_MD + extra)
    here = Path(__file__).resolve()
    (root / ".books").mkdir(exist_ok=True)
    shutil.copy2(here, root / ".books" / "books.py")
    log(root, "init")
    print(f"OK       books created at {root}")
    print(f"         helper copied to {root / '.books' / 'books.py'}; use that copy from now on")
    ignored = git_ignored(root)
    print(("OK       " if ignored else "WARNING  ") + ("the books folder is ignored by Git" if ignored else
          "the books folder is NOT ignored by Git: keep it under output/ or add it to .gitignore"))
    return 0


def company_from_rules(pack: Path) -> dict:
    house = next(pack.rglob("Company and house rules.md"), None)
    out = dict(name="", short="", vat_id="", country="DE")
    if house:
        text = house.read_text()
        m = re.search(r"\| Legal name \| ([^|(]+)", text)
        if m:
            out["name"] = m.group(1).strip()
            out["short"] = out["name"].replace(" GmbH", "").strip()
        m = re.search(r"\| VAT ID \| (\w+)", text)
        if m:
            out["vat_id"] = m.group(1)
        m = re.search(r"\| Bank \| ([^,|(]+)", text)
        if m:
            out["bank"] = m.group(1).strip()
    return out


def log(root: Path, text: str) -> None:
    with (root / "log.md").open("a") as f:
        f.write(f"- {dt.datetime.now().isoformat(timespec='seconds')} {text}\n")


# ---------------------------------------------------------------------------------- collect

class TextOf(html.parser.HTMLParser):
    BLOCK = {"p", "br", "tr", "div", "h1", "h2", "h3", "li", "table", "title"}

    def __init__(self):
        super().__init__()
        self.parts: list[str] = []
        self.skip = 0

    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style"):
            self.skip += 1
        if tag in self.BLOCK:
            self.parts.append("\n")
        if tag == "td":
            self.parts.append("  ")

    def handle_endtag(self, tag):
        if tag in ("script", "style"):
            self.skip -= 1

    def handle_data(self, data):
        if not self.skip:
            self.parts.append(data)


def html_text(raw: str) -> str:
    p = TextOf()
    p.feed(raw)
    lines = [re.sub(r"[ \t\xa0]+", " ", line).strip() for line in "".join(p.parts).splitlines()]
    return "\n".join(line for line in lines if line)


def pdf_text(data: bytes) -> tuple[str, int]:
    try:
        from pypdf import PdfReader
    except ImportError:
        return "(pypdf is not installed: run `poetry install`)", 0
    reader = PdfReader(io.BytesIO(data))
    out = []
    for i, page in enumerate(reader.pages, 1):
        text = (page.extract_text() or "").strip()
        if len(reader.pages) > 1:
            out.append(f"[page {i}]")
        if text:
            out.append(text)
    return "\n".join(out), len(reader.pages)


def xlsx_text(data: bytes) -> tuple[str, str]:
    try:
        import openpyxl
    except ImportError:
        return "(openpyxl is not installed: run `poetry install`)", ""
    wb = openpyxl.load_workbook(io.BytesIO(data))
    values = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    out, note = [], ""
    for ws in wb.worksheets:
        out.append(f"[sheet {ws.title}]")
        vs = values[ws.title]
        for row in ws.iter_rows():
            cells = []
            for c in row:
                if c.value is None:
                    continue
                v = c.value
                if isinstance(v, str) and v.startswith("="):
                    cached = vs[c.coordinate].value
                    v = f"{v} (stored value: {'none' if cached is None else cached})"
                    if cached is None:
                        note = "formula without a stored value"
                cells.append(f"{c.coordinate}: {v}")
            if cells:
                out.append(" | ".join(cells))
    return "\n".join(out), note


def decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", "replace")


def email_text(msg) -> str:
    head = [f"From: {msg.get('From', '')}", f"To: {msg.get('To', '')}"]
    if msg.get("Cc"):
        head.append(f"Cc: {msg.get('Cc')}")
    head += [f"Date: {msg.get('Date', '')}", f"Subject: {msg.get('Subject', '')}"]
    body = msg.get_body(("plain",))
    if body is not None:
        text = body.get_content()
    else:
        html_part = msg.get_body(("html",))
        text = html_text(html_part.get_content()) if html_part is not None else ""
    names = [p.get_filename() for p in msg.iter_attachments() if p.get_filename()]
    if names:
        head.append("Attachments: " + ", ".join(names))
    return "\n".join(head) + "\n\n" + text.strip()


def sources_in(pack: Path, extra: list[str]) -> list[tuple[str, str, bytes]]:
    """(source label, file name, bytes) for every document, including email attachments."""
    found = []
    roots = [pack] + [Path(e) for e in extra]
    for base in roots:
        for p in sorted(base.rglob("*")):
            rel = p.relative_to(base)
            if not p.is_file() or any(part in SKIP_NAMES or part.startswith(".") for part in rel.parts):
                continue
            if p.name in ("expected.json", "evaluation.json") or SECRET_NAME.search(p.name) or p.suffix.lower() == ".mbox":
                continue
            if rel.parts[0] == "stripe" or p.name == "labels.json":
                continue  # Stripe is read with `books.py stripe`; labels are metadata
            data = p.read_bytes()
            label = str(rel) if base == pack else str(p)
            found.append((label, p.name, data))
            if p.suffix.lower() == ".eml":
                msg = email.message_from_bytes(data, policy=policy.default)
                for part in msg.iter_attachments():
                    name = part.get_filename()
                    if name:
                        found.append((f"{label} > {name}", name, part.get_content() if isinstance(part.get_content(), bytes)
                                      else part.get_payload(decode=True)))
    return found


def cmd_collect(args) -> int:
    root, cfg = load_books(args.books)
    pack = Path(cfg["pack"])
    docs: dict[str, dict] = {}
    for label, name, data in sources_in(pack, args.add or []):
        sha = hashlib.sha256(data).hexdigest()
        doc_id = sha[:10]
        if doc_id in docs:
            docs[doc_id]["sources"].append(label)
            continue
        suffix = Path(name).suffix.lower()
        docs[doc_id] = dict(id=doc_id, name=name, suffix=suffix, sha=sha, sources=[label], data=data)
    rows = []
    for d in sorted(docs.values(), key=lambda d: d["sources"][0]):
        inbox = root / "inbox" / f"{d['id']}{d['suffix'] or '.bin'}"
        if not inbox.exists():
            inbox.write_bytes(d["data"])
        kind, pages, note, text = d["suffix"].lstrip(".") or "file", "", "", ""
        try:
            if d["suffix"] == ".pdf":
                text, n = pdf_text(d["data"])
                pages = n
                if not text.replace("[page", "").strip() or len(re.sub(r"\[page \d+\]", "", text).strip()) < 20:
                    note = "no text layer: read it as an image"
                    text = "(no text layer: open the file in inbox/ and read it as an image)"
            elif d["suffix"] in IMAGE:
                note = "image: read it as an image"
                text = "(image: open the file in inbox/ and read it; note how sure you are)"
            elif d["suffix"] == ".eml":
                kind = "email"
                text = email_text(email.message_from_bytes(d["data"], policy=policy.default))
            elif d["suffix"] in (".html", ".htm"):
                text = html_text(decode(d["data"]))
            elif d["suffix"] == ".xlsx":
                text, note = xlsx_text(d["data"])
            elif d["suffix"] in (".csv", ".md", ".txt", ".json"):
                text = decode(d["data"])
            else:
                note = "unknown format"
        except Exception as e:  # a broken file is a finding, not a crash
            note = f"could not read: {e.__class__.__name__}"
        lowered = (d["name"] + " " + " ".join(d["sources"])).lower()
        if d["suffix"] == ".csv" or "kontoauszug" in lowered or "company/" in lowered or "house rules" in lowered:
            need = "no"      # read by the scripts (bank statements, customer and supplier lists, house rules)
        elif kind == "email" or d["suffix"] in (".md", ".txt") or "contract" in lowered or "agreement" in lowered or "lease" in lowered:
            need = "read"    # read it; write a record only if it is a receipt itself or holds a finding
        else:
            need = "yes"
        lines = text.splitlines()
        (root / "text" / f"{d['id']}.txt").write_text("\n".join(f"{i:>4}  {line}" for i, line in enumerate(lines, 1)) + "\n")
        rows.append([d["id"], d["name"], kind, pages, len(lines), "yes" if "image" in note else "no", need, note, len(d["sources"]),
                     " | ".join(d["sources"])])
    write_csv(root / "inventory.csv", ["id", "name", "kind", "pages", "lines", "read_as_image", "needs_record", "note", "copies", "sources"], rows)
    log(root, f"collect: {len(rows)} documents")
    images = sum(1 for r in rows if r[5] == "yes")
    copies = sum(1 for r in rows if r[8] > 1)
    need = sum(1 for r in rows if r[6] == "yes")
    read = sum(1 for r in rows if r[6] == "read")
    print(f"OK       {len(rows)} documents in inventory.csv: {need} need a record, {read} to read (emails, notes, contracts), "
          f"{len(rows) - need - read} are read by the scripts")
    print(f"         {images} must be read as images; {copies} were found more than once")
    print("         next: one record per document, see the extract skill. `books.py todo` lists what is missing.")
    return 0


def cmd_lines(args) -> int:
    root, _ = load_books(args.books)
    p = root / "text" / f"{args.id}.txt"
    if not p.exists():
        print(f"PROBLEM  no text for {args.id}; run collect first")
        return 1
    print(p.read_text())
    return 0


def cmd_todo(args) -> int:
    root, _ = load_books(args.books)
    inv = read_csv(root / "inventory.csv")
    done = {p.stem for p in (root / "records").glob("*.json")}
    need = [r for r in inv if r.get("needs_record") == "yes"]
    missing = [r for r in need if r["id"] not in done]
    for r in missing:
        print(f"TODO     {r['id']}  {r['name']}  ({r['kind']}{', read as image' if r['read_as_image'] == 'yes' else ''})")
    to_read = [r for r in inv if r.get("needs_record") == "read"]
    print(f"{len(need) - len(missing)} of {len(need)} documents that need a record have one.")
    print(f"Also read these {len(to_read)} emails, notes and contracts; write a record only for a receipt or a finding:")
    for r in to_read:
        mark = "done" if r["id"] in done else "    "
        print(f"  {mark}  {r['id']}  {r['name']}")
    return 0


TEMPLATE = {
    "doc_id": "", "type": "bill", "party": "", "number": "", "date": "", "due": "", "currency": "EUR",
    "lines": [{"desc": "", "qty": "1", "unit": "0.00", "amount": "0.00"}],
    "net": "", "vat_rate": "", "vat": "", "gross": "", "stated_gross": "",
    "billed_to": "", "customer_vat_id": "", "iban": "", "refers_to": "", "read_from_image": False,
    "sources": {"number": [0, "exact words from that line"], "stated_gross": [0, ""]},
    "flags": [], "notes": "",
}


def cmd_template(args) -> int:
    t = dict(TEMPLATE, doc_id=args.id)
    print(json.dumps(t, indent=1, ensure_ascii=False))
    return 0


# ----------------------------------------------------------------------------------- verify

def masterdata(cfg: dict) -> tuple[list[dict], list[dict]]:
    pack = Path(cfg["pack"])
    cust = next(pack.rglob("Customers.csv"), None)
    vend = next(pack.rglob("Vendors.csv"), None)
    return (read_csv(cust) if cust else []), (read_csv(vend) if vend else [])


def best(name: str, rows: list[dict]) -> dict | None:
    t = tokens(name)
    scored = [(len(t & tokens(r["name"])), r) for r in rows]
    scored = [s for s in scored if s[0] > 0]
    return max(scored, key=lambda s: s[0])[1] if scored else None


def norm_iban(s: str) -> str:
    return re.sub(r"[^A-Z0-9]", "", (s or "").upper().split("(")[0])


def load_records(root: Path) -> list[dict]:
    out = []
    for p in sorted((root / "records").glob("*.json")):
        if p.name.startswith("_"):
            continue
        try:
            r = json.loads(p.read_text())
        except json.JSONDecodeError as e:
            out.append(dict(doc_id=p.stem, type="broken", _error=str(e)))
            continue
        r.setdefault("doc_id", p.stem)
        out.append(r)
    return out


def text_lines(root: Path, doc_id: str) -> list[str]:
    p = root / "text" / f"{doc_id}.txt"
    if not p.exists():
        return []
    return [re.sub(r"^\s*\d+  ", "", line, count=1) for line in p.read_text().splitlines()]


def squash(s: str) -> str:
    return re.sub(r"\s+", " ", s or "").strip()


def cmd_verify(args) -> int:
    root, cfg = load_books(args.books)
    customers, vendors = masterdata(cfg)
    inv = {r["id"]: r for r in read_csv(root / "inventory.csv")}
    records = load_records(root)
    month = cfg["month"]
    company = cfg.get("company") or {}
    short = (company.get("short") or company.get("name") or "").lower()
    queue, problems, register = [], 0, []

    def flag(code, item, finding, todo, source, by="script"):
        queue.append([code, item, finding, todo, source, by])

    seen_numbers: dict[tuple, str] = {}
    for r in records:
        did = r["doc_id"]
        label = f"{r.get('number') or inv.get(did, {}).get('name', did)}"
        if r.get("type") == "broken":
            print(f"PROBLEM  records/{did}.json is not valid JSON: {r['_error']}")
            problems += 1
            continue
        if did not in inv:
            print(f"PROBLEM  records/{did}.json: no document with id {did} in inventory.csv")
            problems += 1
        lines = text_lines(root, did)
        for field, src in (r.get("sources") or {}).items():
            if not isinstance(src, (list, tuple)) or len(src) != 2:
                print(f"PROBLEM  {did} {field}: source must be [line, \"exact words\"]")
                problems += 1
                continue
            n, quote = src
            if r.get("read_from_image") and not lines[1:]:
                continue
            if not isinstance(n, int) or n < 1 or n > len(lines) or squash(quote) not in squash(lines[n - 1]):
                print(f"PROBLEM  {did} {field}: \"{quote}\" is not on line {n} of text/{did}.txt")
                problems += 1
        # arithmetic, all in Decimal
        net, vat, gross = money(r.get("net")), money(r.get("vat")), money(r.get("gross"))
        stated = money(r.get("stated_gross")) or gross
        rate = money(r.get("vat_rate"))
        line_sum = None
        for i, line in enumerate(r.get("lines") or []):
            q, u, a = money(line.get("qty")), money(line.get("unit")), money(line.get("amount"))
            if q is not None and u is not None and a is not None and (q * u).quantize(CENT, ROUND_HALF_UP) != a:
                flag("LINE_MISMATCH", label, f"line {i + 1}: {q} x {u} is {(q * u).quantize(CENT)}, not {a}", "Ask for a corrected document", did)
            if a is not None:
                line_sum = (line_sum or D("0")) + a
        if line_sum is not None and net is not None and line_sum != net:
            flag("NET_MISMATCH", label, f"lines add up to {line_sum}, net says {net}", "Check the document", did)
        if net is not None and rate is not None and vat is not None and (net * rate / 100).quantize(CENT, ROUND_HALF_UP) != vat:
            flag("VAT_MISMATCH", label, f"{rate}% of {net} is {(net * rate / 100).quantize(CENT, ROUND_HALF_UP)}, not {vat}", "Check the document", did)
        computed = (net + vat) if (net is not None and vat is not None) else gross
        if stated is not None and computed is not None and stated != computed:
            flag("TOTAL_MISMATCH", label, f"states {stated}; {net} + {vat} is {computed}", "Do not pay; ask for a corrected document", did)
        if r.get("type") in ("sales_invoice", "bill", "receipt", "cancellation") and not r.get("currency"):
            flag("MISSING_CURRENCY", label, "no currency on the document", "Ask; never assume EUR", did)
        # duplicates: the same file twice, or the same party and number in two documents
        srcs = (inv.get(did, {}).get("sources") or "").split(" | ")
        in_mail = [x for x in srcs if " > " in x]
        in_files = [x for x in srcs if x and " > " not in x]
        if len(in_mail) > 1 or len(in_files) > 1:
            flag("DUPLICATE", label, f"the same file was found {len(srcs)} times: {' | '.join(srcs)}", "Count it once", did)
        key = (tuple(sorted(tokens(r.get("party", "")))), (r.get("number") or "").strip().upper())
        if key[1] and r.get("type") not in ("email", "note", "contract"):
            if key in seen_numbers:
                flag("DUPLICATE", label, f"{r.get('party')} {r.get('number')} appears in {seen_numbers[key]} and {did}",
                     "Count it once; keep both files", did)
            else:
                seen_numbers[key] = did
        if inv.get(did, {}).get("read_as_image") == "yes" and inv.get(did, {}).get("kind") == "pdf":
            flag("NOT_TEXT", label, "the PDF has no text layer", "Read it as an image; mark values you are unsure of", did)
        if "formula" in inv.get(did, {}).get("note", ""):
            flag("FORMULA_WITHOUT_VALUE", label, "a total is a formula with no stored value", "Calculate it with books.py, never estimate", did)
        typ = r.get("type")
        if typ in ("bill", "receipt"):
            v = best(r.get("party", ""), vendors)
            country = (v or {}).get("country", "DE")
            billed = (r.get("billed_to") or "").lower()
            if (r.get("billed_to") and short and short not in billed) and (country != "DE" or (gross or 0) > 250):
                flag("BILLED_TO_PERSON", label, f"billed to '{r.get('billed_to')}', not to the company", "Ask the supplier to correct it; tell the tax adviser", did)
            no_vat = (vat or D("0")) == 0
            if country != "DE" and no_vat and (country not in EU or not (v or {}).get("vat_id")):
                flag("FOREIGN_NO_VAT", label, f"supplier in {country} without German or EU VAT on the bill", "List it for the tax adviser (reverse charge)", did)
            if r.get("iban") and v and v.get("iban_on_file", "").startswith(("DE", "XX")) and norm_iban(r["iban"]) != norm_iban(v["iban_on_file"]):
                flag("IBAN_CHANGED", label, f"asks for payment to {r['iban']}, the IBAN on file is {v['iban_on_file']}",
                     "Do not pay; call the supplier on the number on file", did)
            if r.get("date") and r["date"][:7] < month:
                flag("OUT_OF_PERIOD", label, f"dated {r['date']}, before {month}", "Leave it out of this month; check if it was booked already", did)
        if typ == "sales_invoice":
            c = best(r.get("party", ""), customers)
            if c and c.get("country") in EU and c.get("country") != "DE" and c.get("type") == "B2B" and (vat or D("0")) == 0 and not r.get("customer_vat_id"):
                flag("RC_WITHOUT_VAT_ID", label, f"reverse charge to {c['country']} without the customer's VAT ID", "Ask the customer for the VAT ID; tell the tax adviser", did)
        for f in r.get("flags") or []:
            flag(f.get("code", "NOTE"), f.get("item") or label, f.get("finding", ""), f.get("what_to_do", ""), did, "agent")
        if typ in ("sales_invoice", "cancellation", "bill", "receipt"):
            register.append([did, typ, r.get("party", ""), r.get("number", ""), r.get("date", ""), r.get("due", ""), r.get("currency", ""),
                             fmt(net), fmt(vat), fmt(computed), fmt(stated), r.get("vat_rate", ""), r.get("refers_to", ""), inv.get(did, {}).get("name", "")])
        if typ == "expense_claim":
            names = {x["name"] for x in inv.values()}
            total = D("0")
            for i, row in enumerate(r.get("rows") or []):
                total += money(row.get("amount")) or D("0")
                rec = row.get("receipt", "")
                if not rec or rec not in names:
                    flag("NO_RECEIPT", f"{label} row {i + 1}", f"{row.get('desc', '')} {row.get('amount', '')}: no receipt found", "Ask for the receipt or a written note", did)
            r["_claim_total"] = fmt(total)
            register.append([did, typ, r.get("party", ""), r.get("number", "") or "expense claim", r.get("date", ""), "", "EUR", "", "", fmt(total), "", "", "", inv.get(did, {}).get("name", "")])
    extra = root / "records" / "_flags.json"
    if extra.exists():
        for f in json.loads(extra.read_text()):
            flag(f["code"], f["item"], f.get("finding", ""), f.get("what_to_do", ""), f.get("source", ""), "agent")
    write_csv(root / "out" / "register.csv", ["doc_id", "type", "party", "number", "date", "due", "currency", "net", "vat", "gross",
                                              "stated_gross", "vat_rate", "refers_to", "file"], register)
    write_csv(root / "out" / "review-queue.csv", QUEUE_HEADER, queue)
    have = {r["doc_id"] for r in records}
    missing = sum(1 for x in inv.values() if x.get("needs_record") == "yes" and x["id"] not in have)
    print(f"{'OK      ' if not problems else 'PROBLEM '} {len(records)} records checked; {problems} quote problem(s); {missing} document(s) still need a record")
    print(f"         {len(register)} rows in out/register.csv; {len(queue)} items in out/review-queue.csv")
    log(root, f"verify: {len(records)} records, {problems} problems, {len(queue)} review items")
    return 1 if problems else 0


def cmd_flag(args) -> int:
    root, _ = load_books(args.books)
    p = root / "records" / "_flags.json"
    items = json.loads(p.read_text()) if p.exists() else []
    items.append(dict(code=args.code.upper(), item=args.item, finding=args.finding, what_to_do=args.todo, source=args.source or ""))
    p.write_text(json.dumps(items, indent=1, ensure_ascii=False) + "\n")
    print(f"OK       {args.code.upper()} {args.item} added; run verify to update the review queue")
    return 0


# -------------------------------------------------------------------------------------- bank

DATE_DE = re.compile(r"^(\d{2})\.(\d{2})\.(\d{4})$")


def de_date(s: str) -> str:
    m = DATE_DE.match((s or "").strip())
    return f"{m.group(3)}-{m.group(2)}-{m.group(1)}" if m else (s or "").strip()


def parse_bank_csv(path: Path) -> tuple[dict, list[dict]]:
    text = decode(path.read_bytes())
    meta, rows, header = {}, [], None
    reader = csv.reader(io.StringIO(text), delimiter=";" if text.count(";") > text.count(",") else ",")
    for line in reader:
        if not line or not any(c.strip() for c in line):
            continue
        if header is None:
            if line[0].strip().lower() in ("buchungstag", "booking date", "date", "datum"):
                header = [c.strip() for c in line]
            elif len(line) >= 2:
                meta[line[0].strip()] = line[1].strip()
            continue
        rows.append(dict(zip(header, line)))
    return meta, rows


def cmd_bank(args) -> int:
    root, cfg = load_books(args.books)
    files = [Path(f) for f in args.files] if args.files else sorted(Path(cfg["pack"]).rglob("*.csv"))
    files = [f for f in files if any(k in f.name.lower() for k in ("umsaetze", "umsätze", "kontoauszug", "bank", "statement"))]
    out, status = [], 0
    for f in files:
        meta, rows = parse_bank_csv(f)
        opening = money(meta.get("Anfangssaldo") or meta.get("Opening balance"))
        closing = money(meta.get("Endsaldo") or meta.get("Closing balance"))
        total = D("0")
        for i, r in enumerate(rows, 1):
            amount = money(r.get("Betrag") or r.get("Amount"))
            total += amount or D("0")
            booked = de_date(r.get("Buchungstag") or r.get("Date", ""))
            out.append([f"B{booked.replace('-', '')}-{i:02d}", booked, de_date(r.get("Valuta", "")), r.get("Auftraggeber/Empfänger", ""),
                        r.get("Verwendungszweck", ""), fmt(amount), r.get("Währung", "EUR"), f.name])
        if opening is not None and closing is not None:
            ok = opening + total == closing
            status |= 0 if ok else 1
            print(f"{'OK      ' if ok else 'PROBLEM '} {f.name}: opening {opening} + lines {total} = {opening + total}; statement says {closing}")
        else:
            print(f"WARNING  {f.name}: no opening or closing balance found; check the total by hand: {total}")
    write_csv(root / "out" / "bank.csv", ["line_id", "date", "value_date", "counterparty", "purpose", "amount", "currency", "file"], out)
    print(f"         {len(out)} bank lines in out/bank.csv")
    log(root, f"bank: {len(out)} lines")
    return status


# ------------------------------------------------------------------------------------ stripe

def env_value(name: str) -> str:
    if os.environ.get(name):
        return os.environ[name].strip()
    for p in (Path(".env"), Path("..") / ".env"):
        if p.exists():
            for line in p.read_text(encoding="utf-8-sig").splitlines():
                m = re.match(rf"\s*(?:export\s+)?{name}\s*=\s*(.*)$", line)
                if m:
                    return m.group(1).strip().strip('"').strip("'")
    return ""


def stripe_get(key: str, path: str, **params) -> list[dict]:
    try:
        import certifi
        ctx = ssl.create_default_context(cafile=certifi.where())
    except ImportError:
        ctx = ssl.create_default_context()
    items, params = [], dict(params, limit=100)
    while True:
        url = f"{STRIPE_API}{path}?{urllib.parse.urlencode(params)}"
        req = urllib.request.Request(url, headers={"Authorization": "Basic " + base64.b64encode(f"{key}:".encode()).decode(),
                                                   "Stripe-Version": STRIPE_VERSION})
        with urllib.request.urlopen(req, context=ctx, timeout=60) as r:
            page = json.loads(r.read())
        items += page["data"]
        if not page.get("has_more"):
            return items
        params["starting_after"] = page["data"][-1]["id"]


def day(epoch) -> str:
    return dt.datetime.fromtimestamp(int(epoch)).strftime("%Y-%m-%d") if epoch else ""


def cents(v) -> str:
    return f"{D(v) / 100:.2f}"


def cmd_stripe(args) -> int:
    root, cfg = load_books(args.books)
    out = root / "out" / "stripe"
    out.mkdir(parents=True, exist_ok=True)
    key = "" if args.snapshot else env_value("STRIPE_API_KEY")
    if key and ("_live_" in key or "_test_" not in key):
        print("PROBLEM  STRIPE_API_KEY is not a sandbox key. The workshop reads Stripe only with a read-only sandbox key.")
        return 1
    if not key:
        snap = Path(cfg["pack"]) / "stripe" / "snapshot"
        if not snap.exists():
            print("PROBLEM  no STRIPE_API_KEY and no snapshot: ask the facilitator for the read-only sandbox key")
            return 1
        for f in snap.glob("*.csv"):
            shutil.copy2(f, out / f.name)
        print(f"OK       Stripe read from the snapshot {snap} (no key set): {', '.join(sorted(p.name for p in out.glob('*.csv')))}")
        log(root, "stripe: snapshot")
        return 0
    try:
        txns = stripe_get(key, "/balance_transactions")
        invoices = stripe_get(key, "/invoices")
        charges = stripe_get(key, "/charges")
        payouts = stripe_get(key, "/payouts")
    except urllib.error.HTTPError as e:
        print(f"PROBLEM  Stripe said {e.code}: {e.read().decode()[:300]}")
        print("         The key needs read access to balance transactions, invoices, charges and payouts.")
        return 1
    write_csv(out / "balance_transactions.csv", ["date", "id", "type", "amount_eur", "fee_eur", "net_eur", "description", "source", "available_on"],
              [[day(t["created"]), t["id"], t["type"], cents(t["amount"]), cents(t["fee"]), cents(t["net"]), t.get("description") or "",
                t.get("source") or "", day(t["available_on"])] for t in sorted(txns, key=lambda t: t["created"]) if t["currency"] == "eur"])
    write_csv(out / "invoices.csv", ["number", "id", "customer_name", "status", "total_eur", "amount_paid_eur", "amount_remaining_eur", "due_date",
                                     "attempt_count", "hosted_invoice_url"],
              [[i.get("number"), i["id"], i.get("customer_name"), i["status"], cents(i["total"]), cents(i.get("amount_paid", 0)),
                cents(i.get("amount_remaining", 0)), day(i.get("due_date")), i.get("attempt_count", 0), i.get("hosted_invoice_url") or ""]
               for i in sorted(invoices, key=lambda i: i.get("number") or "")])
    write_csv(out / "charges.csv", ["date", "id", "status", "amount_eur", "amount_refunded_eur", "description", "customer", "payment_intent",
                                    "failure_code", "invoice_number"],
              [[day(c["created"]), c["id"], c["status"], cents(c["amount"]), cents(c.get("amount_refunded", 0)), c.get("description") or "",
                c.get("customer") or "", c.get("payment_intent") or "", c.get("failure_code") or "", (c.get("metadata") or {}).get("invoice_number", "")]
               for c in sorted(charges, key=lambda c: c["created"])])
    write_csv(out / "payouts.csv", ["created", "id", "amount_eur", "status", "arrival_date", "statement_descriptor"],
              [[day(p["created"]), p["id"], cents(p["amount"]), p["status"], day(p["arrival_date"]), p.get("statement_descriptor") or ""] for p in payouts])
    print(f"OK       read from the Stripe sandbox: {len(invoices)} invoices, {len(charges)} charges, {len(txns)} balance transactions, {len(payouts)} payouts")
    log(root, "stripe: api")
    return 0


# --------------------------------------------------------------------------------- reconcile

def cmd_reconcile(args) -> int:
    root, cfg = load_books(args.books)
    customers, vendors = masterdata(cfg)
    reg = read_csv(root / "out" / "register.csv")
    bank = read_csv(root / "out" / "bank.csv")
    if not reg or not bank:
        print("PROBLEM  run verify and bank first (out/register.csv and out/bank.csv are needed)")
        return 1
    sinv = read_csv(root / "out" / "stripe" / "invoices.csv")
    charges = read_csv(root / "out" / "stripe" / "charges.csv")
    txns = read_csv(root / "out" / "stripe" / "balance_transactions.csv")
    payouts = read_csv(root / "out" / "stripe" / "payouts.csv")
    balance = read_csv(root / "out" / "stripe" / "balance.csv")
    payout_done = False
    queue = read_csv(root / "out" / "review-queue.csv")
    queue = [q for q in queue if q.get("by") != "reconcile"]
    month_end = cfg["month"] + "-31"
    scenario = cfg["scenario_date"]
    bank_tokens = tokens((cfg.get("company") or {}).get("bank", "")) - {"bank"}

    def add(code, item, finding, todo, source):
        queue.append(dict(code=code, item=item, finding=finding, what_to_do=todo, source=source, by="reconcile"))

    docs = {r["number"].upper(): r for r in reg if r.get("number")}
    numbers = sorted(docs, key=len, reverse=True) + [i["number"].upper() for i in sinv if i.get("number")]
    status: dict[str, dict] = {}
    matches = []
    used = set()
    for b in bank:
        amount = money(b["amount"]) or D("0")
        text = f"{b['counterparty']} {b['purpose']}".upper()
        how, target, note = "", "", ""
        if "STRIPE" in text:
            po = next((p for p in payouts if p["id"].upper() in text), None) or next((p for p in payouts if money(p["amount_eur"]) == amount), None)
            if po:
                how, target = "stripe payout", po["id"]
                if not payout_done:
                    payout_done = True
                    gross = sum((money(t["amount_eur"]) for t in txns if t["type"] in ("charge", "payment")), D("0"))
                    fees = sum((money(t["fee_eur"]) for t in txns if t["type"] in ("charge", "payment")), D("0"))
                    refunds = -sum((money(t["amount_eur"]) for t in txns if t["type"] in ("refund", "payment_refund")), D("0"))
                    disputes = -sum((money(t["net_eur"]) for t in txns if t["type"] == "adjustment"), D("0"))
                    net = gross - fees - refunds - disputes
                    paid = sum((money(p["amount_eur"]) for p in payouts), D("0"))
                    rest = sum((money(r.get("available_eur")) or D("0")) + (money(r.get("pending_eur")) or D("0")) for r in balance)
                    note = (f"charged {gross} - fees {fees} - refunds {refunds} - disputes {disputes} = {net}; "
                            f"paid out {paid} in {len(payouts)} payout(s), {rest} still in Stripe")
                    add("PAYOUT_IS_NOT_REVENUE", "STRIPE-PAYOUT", note, "Book the payments, fees, refunds and disputes separately", b["line_id"])
                    if disputes:
                        add("DISPUTE", "STRIPE-DISPUTE", f"a disputed payment took back {disputes} including the dispute fee",
                            "Decide whether to answer the dispute in the Stripe Dashboard", b["line_id"])
                    if refunds:
                        add("REFUND", "STRIPE-REFUNDS", f"refunds of {refunds} reduce the payouts", "Match each refund to its credit note or cancellation", b["line_id"])
                    if net != paid + rest:
                        add("PAYOUT_MISMATCH", "STRIPE-PAYOUT", note, "Check for payments from another period", b["line_id"])
        if not how:
            found = [n for n in numbers if re.search(rf"(?<![A-Z0-9-]){re.escape(n)}(?![A-Z0-9])", text)]
            if found:
                n = found[0]
                how, target = "reference", n
                doc = docs.get(n)
                sdoc = next((i for i in sinv if (i.get("number") or "").upper() == n), None)
                status[n] = dict(state="paid", date=b["date"], channel="bank", line=b["line_id"])
                if doc:
                    gross = money(doc["gross"])
                    if doc["currency"] and doc["currency"] != "EUR":
                        add("CURRENCY", n, f"{doc['currency']} {doc['gross']} arrived as EUR {amount}", "Record the EUR amount; bank fees separately", b["line_id"])
                    elif gross is not None and abs(gross) != abs(amount):
                        add("AMOUNT_DIFFERS", n, f"invoice {gross}, bank {amount}", "Check for a partial payment or a fee", b["line_id"])
                    if doc["type"] == "sales_invoice":
                        cust = best(doc["party"], customers) or {}
                        if not tokens(b["counterparty"]) <= tokens(doc["party"]):
                            add("PAYER_DIFFERS", n, f"paid by {b['counterparty']}, invoice to {doc['party']}", "Matched by number; note the payer", b["line_id"])
                        if b["date"] > month_end and doc["date"] <= month_end:
                            add("PAID_AFTER_MONTH_END", n, f"paid on {b['date']}", "Paid, after the month end", b["line_id"])
                        del cust
                if sdoc and sdoc["status"] == "open":
                    add("PAID_BY_BANK_NOT_STRIPE", n, f"Stripe shows {n} open, the bank shows {amount} on {b['date']}",
                        "Mark it paid outside Stripe after a person approves", b["line_id"])
        if not how and bank_tokens and tokens(b["counterparty"]) & bank_tokens:
            how, target = "bank fee", "no document needed"
        if not how:
            for v in vendors:
                ref = re.findall(r"mandate (\S+)", v.get("pays_by", ""))
                if ref and ref[0].upper() in text:
                    how, target = "direct debit", f"{v['name']} (contract)"
        if not how:
            cands = []
            for r in reg:
                if r["doc_id"] in used or r["type"] not in ("bill", "receipt", "sales_invoice"):
                    continue
                if tokens(r["party"]) & tokens(text) and money(r["gross"]) is not None:
                    same = abs(money(r["gross"])) == abs(amount) or abs(money(r.get("stated_gross") or r["gross"])) == abs(amount)
                    cands.append((same, r))
            exact = [r for same, r in cands if same]
            pick = exact[0] if exact else (cands[0][1] if cands else None)
            if pick:
                how, target = ("name and amount" if exact else "name, other amount"), pick["number"]
                used.add(pick["doc_id"])
                status[pick["number"].upper()] = dict(state="paid", date=b["date"], channel="bank", line=b["line_id"])
                if not exact:
                    add("CURRENCY" if re.search(r"USD|GBP|CHF|KURS", text) else "AMOUNT_DIFFERS", pick["number"],
                        f"document {pick['gross']} {pick['currency']}, bank {amount}", "Record the EUR amount from the bank", b["line_id"])
        if not how:
            add("NO_DOCUMENT", b["line_id"], f"{b['date']} {b['counterparty']} {b['purpose']} {amount}: no document", "Ask who paid and for the receipt", b["line_id"])
        matches.append([b["line_id"], b["date"], b["amount"], b["counterparty"], b["purpose"], target, how, note])
    # payments inside Stripe, by invoice number
    for i in sinv:
        n = (i.get("number") or "").upper()
        if i["status"] == "paid" and n not in status:
            status[n] = dict(state="paid", date="", channel="stripe", line="")
    for c in charges:
        n = (c.get("invoice_number") or "").upper()
        if not n:
            hit = [x for x in numbers if x in (c.get("description") or "").upper()]
            n = hit[0] if hit else ""
        if c["status"] == "succeeded" and n and n not in status:
            status[n] = dict(state="paid", date=c["date"], channel="stripe (payment link)", line=c["id"])
            add("PAID_VIA_STRIPE", n, f"paid by card through Stripe on {c['date']}", "Count it once: it is inside the Stripe payout", c["id"])
    # open items: what is still open, and what is open but must not simply be paid
    claimed = set()
    for rec in load_records(root):
        if rec.get("type") == "expense_claim":
            claimed |= {row.get("receipt", "") for row in rec.get("rows") or [] if row.get("receipt")}
    held = {q["source"]: q["code"] for q in queue if q["code"] in ("IBAN_CHANGED", "TOTAL_MISMATCH", "CEO_FRAUD")}
    month_start = cfg["month"] + "-01"
    rows = []
    for r in reg:
        n = r["number"].upper()
        s = status.get(n)
        if r["type"] == "cancellation" or any(x["refers_to"].upper() == n and x["type"] == "cancellation" for x in reg):
            state = "cancelled"
        elif r["type"] in ("bill", "receipt") and r.get("file") in claimed:
            state = "in expense claim"
        elif r["type"] in ("bill", "receipt") and held.get(r["doc_id"]) == "IBAN_CHANGED":
            state = "blocked: changed bank details"
        elif s:
            state = s["state"]
        elif r["type"] in ("bill", "receipt") and held.get(r["doc_id"]) == "TOTAL_MISMATCH":
            state = "on hold: wrong total"
        elif r["date"] and r["date"] < month_start:
            state = "earlier month"
        else:
            state = "open"
            if r["due"] and r["due"] < scenario:
                code = "OVERDUE" if r["type"] == "sales_invoice" else "PAYABLE_OVERDUE"
                add(code, r["number"], f"due {r['due']}, not paid on {scenario}", "Draft a reminder; a person sends it" if code == "OVERDUE" else "Pay after approval", r["doc_id"])
        rows.append([r["doc_id"], r["type"], r["party"], r["number"], r["date"], r["due"], r["currency"], r["gross"], state,
                     (s or {}).get("date", ""), (s or {}).get("channel", ""), (s or {}).get("line", "")])
    for i in sinv:
        n = (i.get("number") or "").upper()
        if n and n not in {r["number"].upper() for r in reg}:
            s = status.get(n)
            rows.append(["stripe", "sales_invoice", i.get("customer_name", ""), i["number"], "", i.get("due_date", ""), "EUR", i["total_eur"],
                         (s or {}).get("state", i["status"]), (s or {}).get("date", ""), (s or {}).get("channel", "stripe"), (s or {}).get("line", i["id"])])
    write_csv(root / "out" / "status.csv", ["doc_id", "type", "party", "number", "date", "due", "currency", "gross", "state", "paid", "channel", "evidence"], rows)
    write_csv(root / "out" / "matches.csv", ["bank_line", "date", "amount", "counterparty", "purpose", "matched_to", "how", "note"], matches)
    write_csv(root / "out" / "review-queue.csv", QUEUE_HEADER, [[q.get(k, "") for k in QUEUE_HEADER] for q in queue])
    unmatched = sum(1 for m in matches if not m[6])
    print(f"OK       {len(matches)} bank lines: {len(matches) - unmatched} explained, {unmatched} without a document")
    print(f"         out/status.csv: {sum(1 for r in rows if r[8] == 'open')} open, {sum(1 for r in rows if r[8] == 'paid')} paid, "
          f"{sum(1 for r in rows if r[8] == 'cancelled')} cancelled; {len(queue)} items in out/review-queue.csv")
    log(root, f"reconcile: {len(matches)} bank lines, {unmatched} without document")
    return 0


# ----------------------------------------------------------------------------------- package

def cmd_package(args) -> int:
    root, cfg = load_books(args.books)
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        print("PROBLEM  openpyxl is not installed: run `poetry install`")
        return 1
    sheets = [("Status", root / "out" / "status.csv"), ("Register", root / "out" / "register.csv"), ("Bank", root / "out" / "matches.csv"),
              ("Stripe", root / "out" / "stripe" / "balance_transactions.csv"), ("Review notes", root / "out" / "review-queue.csv")]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Summary"
    status = read_csv(root / "out" / "status.csv")
    queue = read_csv(root / "out" / "review-queue.csv")
    bank = read_csv(root / "out" / "bank.csv")
    open_in = [s for s in status if s["state"] == "open" and s["type"] == "sales_invoice"]
    open_out = [s for s in status if s["state"] == "open" and s["type"] in ("bill", "receipt")]
    held = [s for s in status if s["state"].startswith(("blocked", "on hold"))]
    claim = [s for s in status if s["type"] == "expense_claim"]
    summary = [
        ["Month", cfg["month"]], ["Prepared on", cfg["scenario_date"]], ["Company", (cfg.get("company") or {}).get("name", "")],
        ["Bank lines", len(bank)], ["Bank lines without a document", sum(1 for q in queue if q["code"] == "NO_DOCUMENT")],
        ["Customer invoices still open", len(open_in)], ["Open customer invoices, total", fmt(sum((money(s["gross"]) or D("0") for s in open_in), D("0")))],
        ["Bills still open", len(open_out)], ["Items for review", len(queue)],
        ["Status", "DRAFT prepared by an AI agent. A person checks it before it goes to the tax adviser."],
    ]
    for row in summary:
        ws.append(row)
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 70
    for title, path in sheets:
        rows = read_csv(path)
        sh = wb.create_sheet(title)
        if not rows:
            sh.append(["(empty: run the step that writes " + path.name + ")"])
            continue
        sh.append(list(rows[0].keys()))
        for cell in sh[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="E8EEF8")
        for r in rows:
            sh.append([(str(v).lstrip("=+-@") if isinstance(v, str) and v[:1] in "=+@" else v) for v in r.values()])
        for col in sh.columns:
            sh.column_dimensions[col[0].column_letter].width = min(60, max(10, max(len(str(c.value or "")) for c in col) + 2))
    name = f"month-end-{cfg['month']}.xlsx"
    wb.save(root / "out" / name)
    lines = [f"# DRAFT: {cfg['month']} package for the tax adviser", "",
             f"Prepared on {cfg['scenario_date']} by an AI agent; checked by: ________ (a person signs before it is sent).", "",
             "## Open customer invoices"] + [f"- {s['number']} {s['party']}: {s['gross']} {s['currency']}, due {s['due']}" for s in open_in] + [
             "", "## Bills still to pay (after approval)"] + [f"- {s['number']} {s['party']}: {s['gross']} {s['currency']}, due {s['due']}" for s in open_out] + [
             "", "## Do not pay yet"] + [f"- {s['number']} {s['party']}: {s['gross']} {s['currency']}: {s['state']}" for s in held] + [
             "", "## Expense claims (need approval)"] + [f"- {s['party']}: {s['gross']} EUR" for s in claim] + [
             "", "## Questions and findings"] + [f"- **{q['code']}** {q['item']}: {q['finding']} → {q['what_to_do']}" for q in queue]
    (root / "out" / "summary.md").write_text("\n".join(lines) + "\n")
    print(f"OK       out/{name} ({len(sheets) + 1} sheets) and out/summary.md, both marked DRAFT")
    log(root, "package")
    return 0


# ------------------------------------------------------------------------------------- grade

def cmd_grade(args) -> int:
    root, _ = load_books(args.books)
    key = json.loads(Path(args.key).read_text())
    queue = read_csv(root / "out" / "review-queue.csv")
    status = read_csv(root / "out" / "status.csv")
    found, partly, missed = [], [], []
    def about(q, words):
        text = (q["item"] + " " + q["finding"] + " " + q["source"]).upper()
        return any(w.upper() in text for w in words)
    for f in key.get("flags", []):
        hits = [q for q in queue if about(q, f.get("match") or [f["item"]])]
        if any(q["code"].upper() == f["code"] for q in hits):
            found.append(f)
        elif hits:
            partly.append((f, hits[0]["code"]))
        else:
            missed.append(f)
    wrong = []
    for n in key.get("must_not_flag", []):
        bad = [q for q in queue if about(q, n.get("match") or [n["item"]]) and q["code"] in n.get("wrong_codes", [])]
        if bad:
            wrong.append((n, bad[0]["code"]))
    states = {s["number"].upper(): s["state"] for s in status}
    expect = {o["number"]: ("paid" if o["status"] == "paid" else "cancelled" if o["status"] in ("cancelled", "cancellation") else "open")
              for o in key.get("outgoing", [])}
    right = [n for n, s in expect.items() if states.get(n.upper()) == s]
    print(f"Review items: {len(found)} of {len(key.get('flags', []))} found with the right code, {len(partly)} found with another code, {len(missed)} missed")
    for f in missed:
        print(f"  MISSED   {f['code']:<24} {f['item']:<18} {f['finding']}")
    for f, code in partly:
        print(f"  PARTLY   {f['code']:<24} {f['item']:<18} you called it {code}")
    for n, code in wrong:
        print(f"  FALSE    {code:<24} {n['item']:<18} {n['why']}")
    print(f"Invoice status: {len(right)} of {len(expect)} September invoices have the right status")
    for n in sorted(set(expect) - set(right)):
        print(f"  WRONG    {n}: expected {expect[n]}, got {states.get(n.upper(), 'nothing')}")
    print("This grades the fictional pack only. Write down what the agent got wrong: it is half of your demo.")
    return 0


# ------------------------------------------------------------------------------------ doctor

def cmd_doctor(args) -> int:
    print(f"OK       Python {sys.version.split()[0]}")
    for mod, why in (("pypdf", "reads PDFs"), ("openpyxl", "reads and writes Excel")):
        try:
            __import__(mod)
            print(f"OK       {mod} ({why})")
        except ImportError:
            print(f"PROBLEM  {mod} missing ({why}): run `poetry install` in the project")
    root = Path(args.books or "output/accountant")
    if (root / "books.json").exists():
        ign = git_ignored(root)
        print(("OK       " if ign else "WARNING  ") + f"{root} {'is' if ign else 'is NOT'} ignored by Git")
    else:
        print(f"NEEDS SETUP  no books at {root} yet: run init")
    key = env_value("STRIPE_API_KEY")
    if not key:
        print("OPTIONAL no STRIPE_API_KEY: Stripe will be read from the pack's snapshot")
    elif "_live_" in key or "_test_" not in key:
        print("PROBLEM  STRIPE_API_KEY is not a sandbox key; remove it from .env")
    else:
        print("OK       STRIPE_API_KEY is a sandbox key (it is never printed)")
    return 0


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    common = argparse.ArgumentParser(add_help=False)
    common.add_argument("--books", default=None, help="the books folder (default output/accountant)")
    sub = ap.add_subparsers(dest="cmd", required=True)
    p = sub.add_parser("init", help="create the books folder and copy this helper into it")
    p.add_argument("--books", default="output/accountant")
    p.add_argument("--pack", default="data/accountant-pack", help="the folder with the documents")
    p.add_argument("--month", default=None, help="the month to close, YYYY-MM (default: from the pack)")
    sub.add_parser("doctor", parents=[common], help="check Python, libraries, Git safety and the Stripe key")
    p = sub.add_parser("collect", parents=[common], help="copy every document into inbox/ and write numbered text")
    p.add_argument("--add", nargs="*", help="more folders with documents")
    sub.add_parser("todo", parents=[common], help="list documents that still need a record")
    p = sub.add_parser("lines", parents=[common], help="print one document with line numbers")
    p.add_argument("id")
    p = sub.add_parser("template", parents=[common], help="print an empty record for one document")
    p.add_argument("id")
    sub.add_parser("verify", parents=[common], help="check quotes and arithmetic; write register and review queue")
    p = sub.add_parser("flag", parents=[common], help="add a review item that is not tied to one document")
    p.add_argument("code"); p.add_argument("item"); p.add_argument("finding"); p.add_argument("todo")
    p.add_argument("--source", default="")
    p = sub.add_parser("bank", parents=[common], help="read bank statement CSV files and check the balances")
    p.add_argument("files", nargs="*")
    p = sub.add_parser("stripe", parents=[common], help="read Stripe (read-only sandbox key) or the pack's snapshot")
    p.add_argument("--snapshot", action="store_true", help="use the snapshot even if a key is set")
    sub.add_parser("reconcile", parents=[common], help="match bank lines, invoices, bills and Stripe")
    sub.add_parser("package", parents=[common], help="write the DRAFT month-end workbook and summary")
    p = sub.add_parser("grade", parents=[common], help="compare the result with the answer key (fictional pack only)")
    p.add_argument("--key", default="data/accountant-pack/answer-key/expected.json")
    args = ap.parse_args()
    return {"init": cmd_init, "doctor": cmd_doctor, "collect": cmd_collect, "todo": cmd_todo, "lines": cmd_lines, "template": cmd_template,
            "verify": cmd_verify, "flag": cmd_flag, "bank": cmd_bank, "stripe": cmd_stripe, "reconcile": cmd_reconcile,
            "package": cmd_package, "grade": cmd_grade}[args.cmd](args)


if __name__ == "__main__":
    sys.exit(main())
